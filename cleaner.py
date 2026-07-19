#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
CLEANER - Инструменты для работы с ИИ-текстами
Единая точка входа для всех функций
"""

import os
import sys
import argparse
import re
from collections import Counter

if sys.stdout.encoding and sys.stdout.encoding.lower() != "utf-8":
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

# Импорт модулей очистки
try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

# Импорт модулей AI
import watermark
import ai_detector
import pattern_analyzer


# ============================================================================
# ЦВЕТНОЙ ВЫВОД
# ============================================================================

class Colors:
    """ANSI цветовые коды"""
    RESET = '\033[0m'
    BOLD = '\033[1m'

    # Основные цвета
    BLACK = '\033[30m'
    RED = '\033[31m'
    GREEN = '\033[32m'
    YELLOW = '\033[33m'
    BLUE = '\033[34m'
    MAGENTA = '\033[35m'
    CYAN = '\033[36m'
    WHITE = '\033[37m'

    # Яркие цвета
    BRIGHT_RED = '\033[91m'
    BRIGHT_GREEN = '\033[92m'
    BRIGHT_YELLOW = '\033[93m'
    BRIGHT_BLUE = '\033[94m'
    BRIGHT_MAGENTA = '\033[95m'
    BRIGHT_CYAN = '\033[96m'
    BRIGHT_WHITE = '\033[97m'

    # Фон
    BG_BLACK = '\033[40m'
    BG_GREEN = '\033[42m'
    BG_BLUE = '\033[44m'


def colored(text, color):
    """Возвращает цветной текст"""
    return f"{color}{text}{Colors.RESET}"


def print_banner():
    """Выводит красивый баннер"""
    banner = f"""
{Colors.BRIGHT_CYAN}{Colors.BOLD}
╔══════════════════════════════════════════════════════════════════╗
║                                                                  ║
║   {Colors.BRIGHT_YELLOW}██████╗ ██╗     ███████╗ █████╗ ███╗   ██╗███████╗██████╗{Colors.BRIGHT_CYAN}   ║
║   {Colors.BRIGHT_YELLOW}██╔════╝██║     ██╔════╝██╔══██╗████╗  ██║██╔════╝██╔══██╗{Colors.BRIGHT_CYAN}  ║
║   {Colors.BRIGHT_YELLOW}██║     ██║     █████╗  ███████║██╔██╗ ██║█████╗  ██████╔╝{Colors.BRIGHT_CYAN}  ║
║   {Colors.BRIGHT_YELLOW}██║     ██║     ██╔══╝  ██╔══██║██║╚██╗██║██╔══╝  ██╔══██╗{Colors.BRIGHT_CYAN}  ║
║   {Colors.BRIGHT_YELLOW}╚██████╗███████╗███████╗██║  ██║██║ ╚████║███████╗██║  ██║{Colors.BRIGHT_CYAN}  ║
║   {Colors.BRIGHT_YELLOW}╚═════╝╚══════╝╚══════╝╚═╝  ╚═╝╚═╝  ╚═══╝╚══════╝╚═╝  ╚═╝{Colors.BRIGHT_CYAN}  ║
║                                                                  ║
║        {Colors.BRIGHT_WHITE}Инструменты для работы с ИИ-текстами{Colors.BRIGHT_CYAN}                 ║
║                                                                  ║
╚══════════════════════════════════════════════════════════════════╝
{Colors.RESET}"""
    print(banner)


def print_menu():
    """Выводит главное меню"""
    print(f"\n{Colors.BRIGHT_WHITE}{Colors.BOLD}🎯 ГЛАВНОЕ МЕНЮ:{Colors.RESET}\n")

    menu_items = [
        ("1", "🧹 Очистка текста", "Удалить Unicode-символы из документа", Colors.BRIGHT_GREEN),
        ("2", "🤖 Детектор ИИ", "Определить, написан ли текст ИИ (0-100%)", Colors.BRIGHT_BLUE),
        ("3", "🔍 Анализ паттернов", "Найти повторяющиеся конструкции", Colors.BRIGHT_MAGENTA),
        ("4", "🔐 Водяной знак", "Добавить/проверить/удалить водяной знак", Colors.BRIGHT_CYAN),
        ("5", "ℹ️  Справка", "Показать подробную информацию", Colors.BRIGHT_YELLOW),
        ("0", "❌ Выход", "Закрыть программу", Colors.BRIGHT_RED),
    ]

    for num, title, desc, color in menu_items:
        print(f"  {color}{Colors.BOLD}[{num}]{Colors.RESET} {colored(title, color + Colors.BOLD)}")
        print(f"      {Colors.WHITE}{desc}{Colors.RESET}")
        print()


def print_section_header(title, icon=""):
    """Выводит заголовок секции"""
    width = 70
    print(f"\n{Colors.BRIGHT_CYAN}{'═' * width}{Colors.RESET}")
    print(f"{Colors.BRIGHT_WHITE}{Colors.BOLD}{icon} {title.center(width - len(icon) - 2)}{Colors.RESET}")
    print(f"{Colors.BRIGHT_CYAN}{'═' * width}{Colors.RESET}\n")


def print_success(message):
    """Выводит сообщение об успехе"""
    print(f"\n{Colors.BRIGHT_GREEN}✓ {message}{Colors.RESET}\n")


def print_error(message):
    """Выводит сообщение об ошибке"""
    print(f"\n{Colors.BRIGHT_RED}✗ {message}{Colors.RESET}\n")


def print_info(message):
    """Выводит информационное сообщение"""
    print(f"{Colors.BRIGHT_BLUE}ℹ {message}{Colors.RESET}")


def print_warning(message):
    """Выводит предупреждение"""
    print(f"{Colors.BRIGHT_YELLOW}⚠ {message}{Colors.RESET}")


# ============================================================================
# ВСПОМОГАТЕЛЬНЫЕ ФУНКЦИИ ДЛЯ РАБОТЫ С ФАЙЛАМИ
# ============================================================================

def read_text_from_file(file_path):
    """Читает текст из файла (.txt или .docx)"""
    if file_path.endswith('.txt'):
        with open(file_path, 'r', encoding='utf-8') as f:
            return f.read()
    elif file_path.endswith('.docx'):
        if not DOCX_AVAILABLE:
            raise ImportError("Для работы с .docx установите: pip install python-docx")
        doc = Document(file_path)
        return '\n'.join([p.text for p in doc.paragraphs])
    else:
        raise ValueError("Неподдерживаемый формат файла. Используйте .txt или .docx")


def write_text_to_file(file_path, text):
    """Записывает текст в файл (.txt или .docx)"""
    if file_path.endswith('.txt'):
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(text)
    elif file_path.endswith('.docx'):
        if not DOCX_AVAILABLE:
            raise ImportError("Для работы с .docx установите: pip install python-docx")
        # Создаем новый документ
        doc = Document()
        for line in text.split('\n'):
            doc.add_paragraph(line)
        doc.save(file_path)
    else:
        raise ValueError("Неподдерживаемый формат файла. Используйте .txt или .docx")


# ============================================================================
# МОДУЛЬ ОЧИСТКИ ТЕКСТА
# ============================================================================

UNICODE_REPLACEMENTS = {
    '\u00A0': '\u0020',
    '\u202F': '\u0020',
    '\u200B': '',
    '\u00AD': '',
    '\u2019': '\u0027',
    '\u2014': '\u2013',
}

SYMBOL_DESCRIPTIONS = {
    '\u00A0': 'Неразрывный пробел',
    '\u202F': 'Узкий неразрывный пробел',
    '\u200B': 'Пробел нулевой ширины',
    '\u00AD': 'Мягкий перенос',
    '\u2019': 'Правый одинарный апостроф',
    '\u2014': 'Длинное тире',
}


def clean_text(text):
    """Очищает текст от специальных символов Unicode"""
    stats = {}
    cleaned_text = text

    for old_char, new_char in UNICODE_REPLACEMENTS.items():
        count = cleaned_text.count(old_char)
        if count > 0:
            stats[old_char] = count
            cleaned_text = cleaned_text.replace(old_char, new_char)

    return cleaned_text, stats


def analyze_word_frequency(text, top_n=200):
    """Анализирует частоту использования слов"""
    words = re.findall(r'\b[а-яёa-z]+\b', text.lower(), re.UNICODE)
    word_counts = Counter(words)
    return word_counts.most_common(top_n)


def print_clean_statistics(stats, original_length, cleaned_length):
    """Выводит статистику очистки"""
    print(f"\n{Colors.BRIGHT_CYAN}{'─' * 70}{Colors.RESET}")
    print(f"{Colors.BRIGHT_WHITE}{Colors.BOLD}📊 СТАТИСТИКА ОЧИСТКИ{Colors.RESET}")
    print(f"{Colors.BRIGHT_CYAN}{'─' * 70}{Colors.RESET}\n")

    if stats:
        print(f"{Colors.BRIGHT_GREEN}Найденные и замененные символы:{Colors.RESET}\n")
        total_replacements = 0

        for char, count in stats.items():
            description = SYMBOL_DESCRIPTIONS.get(char, "Неизвестный символ")
            replacement = UNICODE_REPLACEMENTS.get(char, '')
            replacement_desc = f"→ '{replacement}'" if replacement else "→ удален"

            print(f"  {Colors.BRIGHT_YELLOW}•{Colors.RESET} {description:35s} "
                  f"{Colors.CYAN}(U+{ord(char):04X}){Colors.RESET}: "
                  f"{Colors.BRIGHT_GREEN}{count:5d}{Colors.RESET} {replacement_desc}")
            total_replacements += count

        print(f"\n{Colors.BRIGHT_WHITE}Всего заменено символов: {Colors.BRIGHT_GREEN}{total_replacements}{Colors.RESET}")
    else:
        print_success("Специальных символов не найдено. Текст уже чистый!")

    print(f"\n{Colors.WHITE}Длина исходного текста:   {Colors.CYAN}{original_length:,}{Colors.WHITE} символов")
    print(f"{Colors.WHITE}Длина очищенного текста:  {Colors.GREEN}{cleaned_length:,}{Colors.WHITE} символов")
    print(f"{Colors.WHITE}Разница:                  {Colors.YELLOW}{original_length - cleaned_length:,}{Colors.WHITE} символов{Colors.RESET}")
    print(f"\n{Colors.BRIGHT_CYAN}{'─' * 70}{Colors.RESET}\n")


# ============================================================================
# ИНТЕРАКТИВНЫЕ ФУНКЦИИ
# ============================================================================

def interactive_clean():
    """Интерактивная очистка текста"""
    print_section_header("ОЧИСТКА ТЕКСТА ОТ UNICODE-СИМВОЛОВ", "🧹")

    # Запрос файла
    file_path = input(f"{Colors.BRIGHT_WHITE}📁 Путь к файлу: {Colors.RESET}").strip()

    if not os.path.exists(file_path):
        print_error(f"Файл '{file_path}' не найден!")
        return

    # Опции
    print(f"\n{Colors.BRIGHT_WHITE}Выберите действие:{Colors.RESET}")
    print(f"  {Colors.GREEN}[1]{Colors.RESET} Создать новый файл")
    print(f"  {Colors.YELLOW}[2]{Colors.RESET} Заменить исходный файл")

    choice = input(f"\n{Colors.BRIGHT_WHITE}Ваш выбор: {Colors.RESET}").strip()

    if choice == '2':
        output_path = file_path
    else:
        default_out = file_path.replace('.txt', '_clean.txt').replace('.docx', '_clean.docx')
        output_path = input(f"{Colors.BRIGHT_WHITE}💾 Выходной файл [{default_out}]: {Colors.RESET}").strip()
        if not output_path:
            output_path = default_out

    # Анализ слов
    analyze = input(f"\n{Colors.BRIGHT_WHITE}📊 Провести анализ частоты слов? (y/n) [{Colors.GREEN}n{Colors.RESET}]: {Colors.RESET}").strip().lower()
    analyze_words = analyze == 'y'

    # Обработка
    print(f"\n{Colors.BRIGHT_CYAN}⏳ Обработка...{Colors.RESET}\n")

    try:
        # Читаем
        original_text = read_text_from_file(file_path)

        # Очищаем
        cleaned_text, stats = clean_text(original_text)

        # Сохраняем
        if output_path.endswith('.txt'):
            write_text_to_file(output_path, cleaned_text)
        elif output_path.endswith('.docx'):
            # Для .docx сохраняем структуру документа
            doc = Document(file_path)
            for paragraph in doc.paragraphs:
                cleaned_para, _ = clean_text(paragraph.text)
                paragraph.text = cleaned_para
            doc.save(output_path)

        print_success(f"Файл сохранен: {output_path}")

        # Статистика
        print_clean_statistics(stats, len(original_text), len(cleaned_text))

        # Анализ слов
        if analyze_words:
            print_section_header("АНАЛИЗ ЧАСТОТЫ СЛОВ (ТОП-200)", "📊")
            word_freq = analyze_word_frequency(cleaned_text)
            for i, (word, count) in enumerate(word_freq, 1):
                color = Colors.BRIGHT_GREEN if count > 10 else Colors.GREEN
                print(f"  {Colors.CYAN}{i:3d}.{Colors.RESET} {word:30s} - {colored(f'{count:5d} раз', color)}")
            print()

    except Exception as e:
        print_error(f"Ошибка: {e}")


def interactive_detect():
    """Интерактивная детекция ИИ"""
    print_section_header("ДЕТЕКТОР ИИ-ТЕКСТА", "🤖")

    file_path = input(f"{Colors.BRIGHT_WHITE}📁 Путь к файлу: {Colors.RESET}").strip()

    if not os.path.exists(file_path):
        print_error(f"Файл '{file_path}' не найден!")
        return

    patterns = input(f"\n{Colors.BRIGHT_WHITE}🔍 Также провести анализ паттернов? (y/n) [{Colors.GREEN}y{Colors.RESET}]: {Colors.RESET}").strip().lower()
    analyze_patterns = patterns != 'n'

    print(f"\n{Colors.BRIGHT_CYAN}⏳ Анализ...{Colors.RESET}\n")

    try:
        text = read_text_from_file(file_path)

        # Детекция
        score, metrics = ai_detector.calculate_ai_score(text)
        ai_detector.print_ai_detection_report(score, metrics)

        # Паттерны
        if analyze_patterns:
            pattern_score, pattern_analysis = pattern_analyzer.detect_ai_writing_patterns(text)
            pattern_analyzer.print_pattern_analysis_report(pattern_score, pattern_analysis)

            # Комбинированная оценка
            combined_score = (score + pattern_score) / 2
            print(f"\n{Colors.BRIGHT_CYAN}{'═' * 70}{Colors.RESET}")
            print(f"{Colors.BRIGHT_WHITE}{Colors.BOLD}📊 КОМБИНИРОВАННАЯ ОЦЕНКА: {Colors.BRIGHT_YELLOW}{combined_score:.1f}%{Colors.RESET}")
            print(f"{Colors.BRIGHT_CYAN}{'═' * 70}{Colors.RESET}\n")

    except Exception as e:
        print_error(f"Ошибка: {e}")


def interactive_watermark():
    """Интерактивное меню водяных знаков"""
    print_section_header("ВОДЯНЫЕ ЗНАКИ", "🔐")

    print(f"{Colors.BRIGHT_WHITE}Выберите действие:{Colors.RESET}\n")
    print(f"  {Colors.GREEN}[1]{Colors.RESET} ➕ Добавить водяной знак")
    print(f"  {Colors.BLUE}[2]{Colors.RESET} 🔍 Проверить водяной знак")
    print(f"  {Colors.RED}[3]{Colors.RESET} 🗑️  Удалить водяной знак")
    print(f"  {Colors.YELLOW}[0]{Colors.RESET} ⬅️  Назад")

    choice = input(f"\n{Colors.BRIGHT_WHITE}Ваш выбор: {Colors.RESET}").strip()

    if choice == '1':
        interactive_watermark_add()
    elif choice == '2':
        interactive_watermark_check()
    elif choice == '3':
        interactive_watermark_remove()


def interactive_watermark_add():
    """Добавление водяного знака"""
    print_section_header("ДОБАВЛЕНИЕ ВОДЯНОГО ЗНАКА", "➕")

    file_path = input(f"{Colors.BRIGHT_WHITE}📁 Путь к файлу: {Colors.RESET}").strip()

    if not os.path.exists(file_path):
        print_error(f"Файл '{file_path}' не найден!")
        return

    author = input(f"{Colors.BRIGHT_WHITE}✍️  Имя автора: {Colors.RESET}").strip()
    if not author:
        print_error("Имя автора обязательно!")
        return

    print(f"\n{Colors.BRIGHT_WHITE}Тип водяного знака:{Colors.RESET}")
    print(f"  {Colors.GREEN}[1]{Colors.RESET} Концентрированный (метаданные в начале)")
    print(f"  {Colors.BLUE}[2]{Colors.RESET} Распределенный (по всему тексту, устойчивее)")

    wm_type = input(f"\n{Colors.BRIGHT_WHITE}Ваш выбор [{Colors.GREEN}1{Colors.RESET}]: {Colors.RESET}").strip()
    distributed = wm_type == '2'

    output_path = input(f"\n{Colors.BRIGHT_WHITE}💾 Выходной файл [{file_path}]: {Colors.RESET}").strip()
    if not output_path:
        output_path = file_path

    print(f"\n{Colors.BRIGHT_CYAN}⏳ Добавление водяного знака...{Colors.RESET}\n")

    try:
        text = read_text_from_file(file_path)

        if distributed:
            watermarked_text = watermark.create_distributed_watermark(text, author, density=0.15)
            print_info(f"Тип: {Colors.BLUE}Распределенный{Colors.RESET}")
        else:
            watermarked_text = watermark.embed_watermark(text, author)
            print_info(f"Тип: {Colors.GREEN}Концентрированный{Colors.RESET}")

        write_text_to_file(output_path, watermarked_text)

        print_success(f"Водяной знак добавлен: {output_path}")
        print(f"{Colors.WHITE}Автор: {Colors.BRIGHT_GREEN}{author}{Colors.RESET}")
        print(f"{Colors.WHITE}Добавлено символов: {Colors.BRIGHT_YELLOW}{len(watermarked_text) - len(text)}{Colors.RESET}\n")

    except Exception as e:
        print_error(f"Ошибка: {e}")


def interactive_watermark_check():
    """Проверка водяного знака"""
    print_section_header("ПРОВЕРКА ВОДЯНОГО ЗНАКА", "🔍")

    file_path = input(f"{Colors.BRIGHT_WHITE}📁 Путь к файлу: {Colors.RESET}").strip()

    if not os.path.exists(file_path):
        print_error(f"Файл '{file_path}' не найден!")
        return

    print(f"\n{Colors.BRIGHT_CYAN}⏳ Проверка...{Colors.RESET}\n")

    try:
        text = read_text_from_file(file_path)

        # Проверка концентрированного
        extracted = watermark.extract_watermark(text)

        if extracted:
            print_success("ВОДЯНОЙ ЗНАК ОБНАРУЖЕН!")
            watermark.print_watermark_info(extracted)
        else:
            print_warning("Концентрированный водяной знак не найден")

            # Проверка распределенного
            author = input(f"\n{Colors.BRIGHT_WHITE}Проверить распределенный? Введите имя автора (Enter - пропустить): {Colors.RESET}").strip()
            if author:
                found, confidence = watermark.detect_distributed_watermark(text, author)
                print(f"\n{Colors.WHITE}Автор: {Colors.CYAN}{author}{Colors.RESET}")
                print(f"{Colors.WHITE}Уверенность: {Colors.BRIGHT_YELLOW}{confidence:.1%}{Colors.RESET}")

                if found:
                    print_success(f"Распределенный водяной знак найден!")
                else:
                    print_warning(f"Распределенный водяной знак не найден")

        print()

    except Exception as e:
        print_error(f"Ошибка: {e}")


def interactive_watermark_remove():
    """Удаление водяного знака"""
    print_section_header("УДАЛЕНИЕ ВОДЯНОГО ЗНАКА", "🗑️")

    file_path = input(f"{Colors.BRIGHT_WHITE}📁 Путь к файлу: {Colors.RESET}").strip()

    if not os.path.exists(file_path):
        print_error(f"Файл '{file_path}' не найден!")
        return

    output_path = input(f"{Colors.BRIGHT_WHITE}💾 Выходной файл [{file_path}]: {Colors.RESET}").strip()
    if not output_path:
        output_path = file_path

    print(f"\n{Colors.BRIGHT_CYAN}⏳ Удаление...{Colors.RESET}\n")

    try:
        text = read_text_from_file(file_path)

        cleaned_text = watermark.remove_watermark(text)

        write_text_to_file(output_path, cleaned_text)

        print_success(f"Водяной знак удален: {output_path}")
        print(f"{Colors.WHITE}Удалено символов: {Colors.BRIGHT_YELLOW}{len(text) - len(cleaned_text)}{Colors.RESET}\n")

    except Exception as e:
        print_error(f"Ошибка: {e}")


def show_help():
    """Показать справку"""
    print_section_header("СПРАВКА И ДОКУМЕНТАЦИЯ", "ℹ️")

    help_text = f"""
{Colors.BRIGHT_WHITE}{Colors.BOLD}📚 БЫСТРЫЙ СТАРТ{Colors.RESET}

{Colors.BRIGHT_GREEN}1. Очистка текста{Colors.RESET}
   Удаляет невидимые символы Unicode из документов.
   Поддерживает: .txt, .docx

   {Colors.CYAN}Что удаляется:{Colors.RESET}
   • Неразрывные пробелы (U+00A0, U+202F)
   • Пробелы нулевой ширины (U+200B)
   • Мягкие переносы (U+00AD)
   • Правые апострофы → обычные (')
   • Длинные тире (—) → короткие (–)

{Colors.BRIGHT_BLUE}2. Детектор ИИ{Colors.RESET}
   Определяет, написан ли текст искусственным интеллектом.

   {Colors.CYAN}Анализирует:{Colors.RESET}
   • Однородность длины предложений
   • Характерные фразы ИИ
   • Частоту переходных слов
   • Лексическое разнообразие

   {Colors.GREEN}Работает локально, без интернета!{Colors.RESET}

{Colors.BRIGHT_MAGENTA}3. Анализ паттернов{Colors.RESET}
   Находит повторяющиеся конструкции.

   {Colors.CYAN}Ищет:{Colors.RESET}
   • Повторяющиеся n-граммы (2-5 слов)
   • Шаблонные выражения
   • Чрезмерно используемую лексику

{Colors.BRIGHT_CYAN}4. Водяные знаки{Colors.RESET}
   Защита авторства через невидимые метки.

   {Colors.CYAN}Типы:{Colors.RESET}
   • Концентрированный - метаданные в начале
   • Распределенный - по всему тексту (устойчивее)

{Colors.BRIGHT_WHITE}{Colors.BOLD}💻 КОМАНДНАЯ СТРОКА{Colors.RESET}

Вы также можете использовать прямые команды:

  {Colors.GREEN}python cleaner.py clean input.txt -o output.txt{Colors.RESET}
  {Colors.BLUE}python cleaner.py detect --file document.txt --patterns{Colors.RESET}
  {Colors.CYAN}python cleaner.py watermark-add --file doc.txt --author "Имя"{Colors.RESET}

Справка по команде: {Colors.YELLOW}python cleaner.py <команда> --help{Colors.RESET}
"""
    print(help_text)
    input(f"\n{Colors.BRIGHT_WHITE}Нажмите Enter для продолжения...{Colors.RESET}")


def interactive_mode():
    """Интерактивный режим"""
    while True:
        print_banner()
        print_menu()

        choice = input(f"{Colors.BRIGHT_WHITE}{Colors.BOLD}➤ Ваш выбор: {Colors.RESET}").strip()

        if choice == '1':
            interactive_clean()
        elif choice == '2':
            interactive_detect()
        elif choice == '3':
            file_path = input(f"\n{Colors.BRIGHT_WHITE}📁 Путь к файлу: {Colors.RESET}").strip()
            if os.path.exists(file_path):
                print(f"\n{Colors.BRIGHT_CYAN}⏳ Анализ...{Colors.RESET}\n")
                try:
                    text = read_text_from_file(file_path)
                    score, analysis = pattern_analyzer.detect_ai_writing_patterns(text)
                    pattern_analyzer.print_pattern_analysis_report(score, analysis)
                except Exception as e:
                    print_error(f"Ошибка: {e}")
            else:
                print_error(f"Файл не найден!")
        elif choice == '4':
            interactive_watermark()
        elif choice == '5':
            show_help()
        elif choice == '0':
            print(f"\n{Colors.BRIGHT_GREEN}👋 До свидания!{Colors.RESET}\n")
            sys.exit(0)
        else:
            print_error("Неверный выбор! Введите номер от 0 до 5.")

        input(f"\n{Colors.BRIGHT_WHITE}Нажмите Enter для продолжения...{Colors.RESET}")


# ============================================================================
# CLI РЕЖИМ (для обратной совместимости)
# ============================================================================

def cli_mode():
    """CLI режим с аргументами"""
    parser = argparse.ArgumentParser(
        description='CLEANER - Инструменты для работы с ИИ-текстами',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Примеры:
  python cleaner.py                    # Интерактивный режим
  python cleaner.py clean input.txt -o output.txt
  python cleaner.py detect --file document.txt --patterns
  python cleaner.py watermark-add --file doc.txt --author "Имя"
        """
    )

    subparsers = parser.add_subparsers(dest='command', help='Команда для выполнения')

    # CLEAN
    parser_clean = subparsers.add_parser('clean', help='Очистка текста')
    parser_clean.add_argument('input', help='Входной файл')
    parser_clean.add_argument('-o', '--output', help='Выходной файл')
    parser_clean.add_argument('--inplace', action='store_true', help='Заменить исходный')
    parser_clean.add_argument('--analyze-words', action='store_true', help='Анализ слов')

    # DETECT
    parser_detect = subparsers.add_parser('detect', help='Детекция ИИ')
    parser_detect.add_argument('--file', required=True, help='Файл')
    parser_detect.add_argument('--patterns', action='store_true', help='С паттернами')

    # PATTERNS
    parser_patterns = subparsers.add_parser('patterns', help='Анализ паттернов')
    parser_patterns.add_argument('--file', required=True, help='Файл')

    # WATERMARK-ADD
    parser_wm_add = subparsers.add_parser('watermark-add', help='Добавить водяной знак')
    parser_wm_add.add_argument('--file', required=True, help='Файл')
    parser_wm_add.add_argument('--author', required=True, help='Автор')
    parser_wm_add.add_argument('-o', '--output', help='Выходной файл')
    parser_wm_add.add_argument('--distributed', action='store_true', help='Распределенный')
    parser_wm_add.add_argument('--density', type=float, default=0.15, help='Плотность')

    # WATERMARK-CHECK
    parser_wm_check = subparsers.add_parser('watermark-check', help='Проверить водяной знак')
    parser_wm_check.add_argument('--file', required=True, help='Файл')
    parser_wm_check.add_argument('--author', help='Автор (для распределенного)')

    # WATERMARK-REMOVE
    parser_wm_remove = subparsers.add_parser('watermark-remove', help='Удалить водяной знак')
    parser_wm_remove.add_argument('--file', required=True, help='Файл')
    parser_wm_remove.add_argument('-o', '--output', help='Выходной файл')

    args = parser.parse_args()

    # Если нет команды - интерактивный режим
    if not args.command:
        return None

    return args


def _read_text_file(file_path):
    """Читает .txt или .docx файл, возвращает текст."""
    if file_path.endswith('.docx'):
        if not DOCX_AVAILABLE:
            print_error("Для работы с .docx установите: pip install python-docx")
            return None
        doc = Document(file_path)
        return '\n'.join(p.text for p in doc.paragraphs)
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()


def _write_text_file(file_path, source_path, text):
    """Сохраняет текст в .txt или .docx (для .docx копирует форматирование source_path)."""
    if file_path.endswith('.docx'):
        doc = Document(source_path)
        paragraphs = text.split('\n')
        for paragraph, new_text in zip(doc.paragraphs, paragraphs):
            paragraph.text = new_text
        doc.save(file_path)
    else:
        with open(file_path, 'w', encoding='utf-8') as f:
            f.write(text)


def run_cli_command(args):
    """Выполняет команду, переданную через аргументы командной строки."""
    if args.command == 'clean':
        if not os.path.exists(args.input):
            print_error(f"Файл '{args.input}' не найден!")
            return
        text = _read_text_file(args.input)
        if text is None:
            return

        output_path = args.input if args.inplace else (args.output or args.input.rsplit('.', 1)[0] + '_clean.' + args.input.rsplit('.', 1)[-1])
        cleaned_text, stats = clean_text(text)
        _write_text_file(output_path, args.input, cleaned_text)

        print_success(f"Файл сохранен: {output_path}")
        print_clean_statistics(stats, len(text), len(cleaned_text))

        if args.analyze_words:
            print_section_header("АНАЛИЗ ЧАСТОТЫ СЛОВ (ТОП-200)", "📊")
            for i, (word, count) in enumerate(analyze_word_frequency(cleaned_text), 1):
                print(f"  {i:3d}. {word:30s} - {count:5d} раз")

    elif args.command == 'detect':
        if not os.path.exists(args.file):
            print_error(f"Файл '{args.file}' не найден!")
            return
        text = _read_text_file(args.file)
        if text is None:
            return

        score, metrics = ai_detector.calculate_ai_score(text)
        ai_detector.print_ai_detection_report(score, metrics)

        if args.patterns:
            pattern_score, pattern_analysis = pattern_analyzer.detect_ai_writing_patterns(text)
            pattern_analyzer.print_pattern_analysis_report(pattern_score, pattern_analysis)
            combined_score = (score + pattern_score) / 2
            print(f"\n📊 КОМБИНИРОВАННАЯ ОЦЕНКА: {combined_score:.1f}%\n")

    elif args.command == 'patterns':
        if not os.path.exists(args.file):
            print_error(f"Файл '{args.file}' не найден!")
            return
        text = _read_text_file(args.file)
        if text is None:
            return
        pattern_score, pattern_analysis = pattern_analyzer.detect_ai_writing_patterns(text)
        pattern_analyzer.print_pattern_analysis_report(pattern_score, pattern_analysis)

    elif args.command == 'watermark-add':
        if not os.path.exists(args.file):
            print_error(f"Файл '{args.file}' не найден!")
            return
        text = _read_text_file(args.file)
        if text is None:
            return

        if args.distributed:
            marked_text = watermark.create_distributed_watermark(text, args.author, args.density)
        else:
            marked_text = watermark.embed_watermark(text, args.author)

        output_path = args.output or args.file
        _write_text_file(output_path, args.file, marked_text)
        print_success(f"Водяной знак добавлен: {output_path}")

    elif args.command == 'watermark-check':
        if not os.path.exists(args.file):
            print_error(f"Файл '{args.file}' не найден!")
            return
        text = _read_text_file(args.file)
        if text is None:
            return

        if args.author:
            found, confidence = watermark.detect_distributed_watermark(text, args.author)
            if found:
                print_success(f"Распределенный водяной знак найден (уверенность: {confidence:.1%})")
            else:
                print_warning("Распределенный водяной знак не найден")
        else:
            data = watermark.extract_watermark(text)
            if data:
                watermark.print_watermark_info(data)
            else:
                print_warning("Водяной знак не найден")

    elif args.command == 'watermark-remove':
        if not os.path.exists(args.file):
            print_error(f"Файл '{args.file}' не найден!")
            return
        text = _read_text_file(args.file)
        if text is None:
            return

        cleaned_text = watermark.remove_watermark(text)
        output_path = args.output or args.file
        _write_text_file(output_path, args.file, cleaned_text)
        print_success(f"Водяной знак удален: {output_path}")


# ============================================================================
# ГЛАВНАЯ ФУНКЦИЯ
# ============================================================================

def main():
    """Главная функция"""
    # Проверяем аргументы командной строки
    if len(sys.argv) > 1:
        # CLI режим
        args = cli_mode()
        if args is None:
            print_info("Запустите без аргументов: python cleaner.py")
        else:
            run_cli_command(args)
    else:
        # Интерактивный режим
        try:
            interactive_mode()
        except KeyboardInterrupt:
            print(f"\n\n{Colors.BRIGHT_YELLOW}⚠ Прервано пользователем{Colors.RESET}")
            print(f"{Colors.BRIGHT_GREEN}👋 До свидания!{Colors.RESET}\n")
            sys.exit(0)


if __name__ == '__main__':
    main()
